#!/usr/bin/env python3
"""Generate Excel workbooks: regulatory matrix, GH capability, security, interfaces, master data."""

from __future__ import annotations

from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[2]
RESEARCH = ROOT / "docs" / "research"
SECURITY = ROOT / "docs" / "security"
INTEGRATION = ROOT / "docs" / "integration"
TEMPLATES = ROOT / "templates"
TODAY = date.today().isoformat()

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
WRAP = Alignment(wrap_text=True, vertical="top")


def style_header(ws, cols: int) -> None:
    for c in range(1, cols + 1):
        cell = ws.cell(1, c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = WRAP
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def autosize(ws, max_width: int = 40) -> None:
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        width = min(max_width, max(len(str(c.value or "")) for c in col) + 2)
        ws.column_dimensions[letter].width = max(12, width)


def add_rows(ws, headers: list[str], rows: list[list]) -> None:
    ws.append(headers)
    for r in rows:
        ws.append(r)
    style_header(ws, len(headers))
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=len(headers)):
        for cell in row:
            cell.alignment = WRAP
    autosize(ws)


def regulatory_matrix() -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Traceability"
    headers = [
        "ReqID", "Requirement", "Country", "Authority", "OfficialSource",
        "SourceDate", "Classification", "GHModule", "SupportLevel",
        "Gap", "Action", "Owner", "Status",
    ]
    rows = [
        ["REQ-UAE-001", "HIE connectivity for Dubai facilities (NABIDH)", "UAE", "DHA",
         "NABIDH Interoperability Standards PDF", "2022-10", "Needs legal confirmation",
         "None", "Unsupported", "No HL7/FHIR gateway", "Engage DHA; design interface service", "Compliance", "Open"],
        ["REQ-UAE-002", "Support HL7 v2 for HIE submission", "UAE", "DHA/MOHAP",
         "NABIDH + Riayati standards", "2022", "Recommended (published standards)",
         "None", "Unsupported", "No message engine", "Implement adapter", "Integration", "Open"],
        ["REQ-UAE-003", "Support FHIR APIs where required by HIE", "UAE", "DHA/MOHAP",
         "Riayati Terminology & Coding Standards v2.1", "2022-06-29", "Recommended",
         "None", "Unsupported", "No FHIR server/client", "FHIR gateway design", "Integration", "Open"],
        ["REQ-UAE-004", "C-CDA clinical document exchange", "UAE", "DHA",
         "NABIDH implementation guidelines", "n/a", "Needs legal confirmation",
         "None", "Unsupported", "No C-CDA export", "Map clinical docs", "Integration", "Open"],
        ["REQ-UAE-005", "IHE profiles (XDS/PDQ/etc.) as applicable", "UAE", "MOHAP/DHA",
         "Riayati standards", "2022", "Needs legal confirmation",
         "None", "Unsupported", "No IHE actors", "Confirm required profiles", "Compliance", "Open"],
        ["REQ-UAE-006", "Malaffi participation (Abu Dhabi)", "UAE", "DoH",
         "Malaffi program communications", "2023+", "Needs legal confirmation",
         "None", "Unsupported", "No connector", "Facility onboarding if in scope", "Compliance", "Open"],
        ["REQ-UAE-007", "Riayati / NUMR alignment", "UAE", "MOHAP",
         "MOHAP announcement + Riayati tech docs", "2023+", "Needs legal confirmation",
         "None", "Unsupported", "No connector", "Confirm facility obligations", "Compliance", "Open"],
        ["REQ-UAE-008", "UAE PDPL controls for health/sensitive data", "UAE", "Federal PDPL",
         "Federal Decree-Law No. 45/2021 (confirm text)", "", "Needs legal confirmation",
         "health_crypto (partial)", "Partial", "No DPIA/process pack", "Counsel + DPIA", "Legal", "Open"],
        ["REQ-UAE-009", "Coding standards (ICD/LOINC/SNOMED as specified)", "UAE", "MOHAP/DHA",
         "Riayati coding standards", "2022", "Needs legal confirmation",
         "health_icd10/11", "Partial", "SNOMED/LOINC licensing unknown", "Confirm licenses & maps", "Clinical Informatics", "Open"],
        ["REQ-QA-001", "MoPH facility HIS / e-health obligations", "Qatar", "MoPH",
         "moph.gov.qa — specific mandate TBD", "", "Needs legal confirmation",
         "health*", "Partial", "Mandate set not confirmed", "Formal MoPH inquiry", "Compliance", "Open"],
        ["REQ-QA-002", "PDPPL special-nature health data processing", "Qatar", "NDPO/NCSA",
         "Law No. 13 of 2016", "2016/2017", "Needs legal confirmation",
         "N/A", "Partial", "Permits/DPIA not implemented in product", "Counsel + technical controls", "Legal", "Open"],
        ["REQ-QA-003", "Patient consent & access controls", "Qatar", "MoPH/NDPO",
         "PDPPL + MoPH patient rights (confirm)", "", "Needs legal confirmation",
         "res.group / record rules", "Partial", "No consent registry UI", "Design consent module", "Product", "Open"],
        ["REQ-BOTH-001", "Audit trail for PHI access", "UAE/QA", "Privacy authorities",
         "PDPL/PDPPL principles (confirm)", "", "Recommended / confirm mandatory",
         "Tryton model history", "Partial", "No privacy-centric audit UI", "Enhance logging", "Security", "Open"],
        ["REQ-BOTH-002", "Tenant isolation multi-country", "UAE/QA", "Governance",
         "Internal policy", TODAY, "Recommended",
         "company.company", "Partial", "Single demo company", "Org hierarchy + rules", "Architecture", "In progress"],
        ["REQ-BOTH-003", "No real PII in non-prod", "UAE/QA", "Internal",
         "IST policy", TODAY, "Mandatory (internal)",
         "synthetic_data", "Supported", "Need production gate", "Refuse prod import without override", "Engineering", "In progress"],
        ["REQ-INT-001", "HL7 trademark / FHIR use compliance", "Intl", "HL7",
         "hl7.org IP policy", "", "Recommended",
         "N/A", "N/A", "Docs only", "Follow IP policy", "Legal", "Open"],
        ["REQ-INT-002", "SNOMED CT affiliate license if used", "Intl/UAE/QA", "SNOMED Int'l",
         "snomed.org", "", "Needs legal confirmation",
         "None verified", "Unsupported", "License unknown", "Confirm national membership", "Clinical Informatics", "Open"],
        ["REQ-INT-003", "LOINC attribution license", "Intl", "Regenstrief",
         "loinc.org/license", "", "Recommended if used",
         "None verified", "Unsupported", "Not loaded", "Adopt with attribution", "Clinical Informatics", "Open"],
        ["REQ-INT-004", "ICD commercial redistribution rules", "Intl", "WHO",
         "WHO classifications", "", "Needs legal confirmation",
         "health_icd*", "Partial", "National license unclear", "Confirm with focal point", "Legal", "Open"],
        ["REQ-INT-005", "DICOM if imaging/PACS integrated", "Intl", "NEMA",
         "dicomstandard.org", "", "Recommended if imaging",
         "health_imaging / orthanc", "Partial", "Not wired in React", "Optional Orthanc", "Imaging", "Open"],
    ]
    add_rows(ws, headers, rows)
    meta = wb.create_sheet("Legend")
    add_rows(
        meta,
        ["Classification", "Meaning"],
        [
            ["Mandatory", "Confirmed legal/regulatory obligation — only after authority confirmation"],
            ["Recommended", "Published standard or best practice; confirm applicability"],
            ["Needs legal confirmation", "Do not treat as binding until counsel/authority confirms"],
            ["Internal", "IST project policy"],
        ],
    )
    path = RESEARCH / "regulatory-traceability-matrix.xlsx"
    wb.save(path)
    return path


def capability_matrix() -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Capabilities"
    headers = [
        "Area", "Capability", "GHModule", "UpstreamSupport", "ISTReact",
        "Evidence", "UAERelevance", "QatarRelevance", "GapAction",
    ]
    rows = [
        ["Demographics", "Patient master", "health", "Full", "Live", "models.ts", "High", "High", "National ID rules TBD"],
        ["Outpatient", "Appointments", "health/health_calendar", "Full", "Live", "AppointmentsPage", "High", "High", ""],
        ["Clinical", "Evaluations", "health", "Full", "Live subset", "EvaluationsPage", "High", "High", "Expand forms"],
        ["Pharmacy", "Prescriptions", "health", "Full", "Live", "PharmacyPage", "High", "High", ""],
        ["Lab", "Lab orders/results", "health_lab", "Full", "Live if module on", "LaboratoryPage", "High", "High", "VM verify activation"],
        ["Imaging", "Dx imaging", "health_imaging", "Full", "Live if module on", "ImagingPage", "Medium", "Medium", "Orthanc optional"],
        ["Inpatient", "Beds/admissions", "health_inpatient", "Full", "Live", "BedsPage", "Medium", "Medium", ""],
        ["Nursing", "Rounding/ambulatory", "health_nursing", "Full", "Live if on", "NursingPage", "High", "High", ""],
        ["Surgery", "OR / procedures", "health_surgery", "Full", "Live if on", "ProceduresPage", "Medium", "Medium", ""],
        ["ICU", "Critical care", "health_icu", "Full", "No dedicated UI", "SAO", "Low-Med", "Low-Med", "ModelBrowser/SAO"],
        ["Billing", "Invoicing", "health_accounting", "Full", "Stub", "FinancialPage", "High", "High", "SAO for CoA"],
        ["Insurance", "Plans/policies", "health_insurance", "Full", "None", "SAO", "High", "High", "UI + claims gap"],
        ["Stock", "Inventory", "health_stock", "Full", "List", "StockPage", "High", "High", ""],
        ["Reporting", "Stats/epidemics", "health_reporting", "Full", "None", "SAO", "Medium", "Medium", ""],
        ["Security", "Digital signatures", "health_crypto", "Documented", "None", "docs", "Medium", "Medium", "Plugin testing"],
        ["Interop", "Federation", "health_federation", "Documented", "Unused", "docs.gnuhealth.org", "Low", "Low", "Prefer HIE adapters"],
        ["Coding", "ICD-10", "health_icd10", "Full", "None dedicated", "features.html", "High", "High", "License confirm"],
        ["Coding", "ICD-11", "health_icd11", "Full", "None dedicated", "features.html", "Medium", "Medium", "License confirm"],
        ["Admin", "Users/groups", "Tryton res.*", "Full", "Live", "OpsPages", "High", "High", "Role packs"],
        ["HIE", "NABIDH/Malaffi/Riayati", "N/A", "None", "None", "gap-analysis", "Critical if licensed", "N/A", "External service"],
    ]
    add_rows(ws, headers, rows)
    path = RESEARCH / "gnuhealth-capability-matrix.xlsx"
    wb.save(path)
    return path


def security_matrix() -> Path:
    SECURITY.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Controls"
    headers = [
        "ControlID", "Domain", "Control", "UAERef", "QatarRef",
        "Implementation", "Status", "Evidence", "ResidualRisk",
    ]
    rows = [
        ["CTL-01", "Access", "Unique user accounts", "PDPL principles — confirm", "PDPPL — confirm",
         "res.user", "Partial", "UsersPage", "Shared accounts possible"],
        ["CTL-02", "Access", "Role-based groups", "Confirm", "Confirm", "res.group + ir.model.access", "Partial", "GroupsPage", "No IST role pack yet"],
        ["CTL-03", "Access", "Session security", "Confirm", "Confirm", "Tryton Session; localStorage MVP", "Gap", "ARCHITECTURE.md", "Move to httpOnly/JWT"],
        ["CTL-04", "Data", "Encryption in transit", "Confirm", "Confirm", "HTTPS nginx", "Partial", "HTTPS.md", "Confirm TLS config"],
        ["CTL-05", "Data", "Encryption at rest", "Confirm", "Confirm", "GCP disk — verify", "Needs verification", "VM", "Enable CMEK if required"],
        ["CTL-06", "Privacy", "Synthetic-only non-prod", "Internal", "Internal", "synthetic markers + import gate", "In progress", "synthetic factory", ""],
        ["CTL-07", "Privacy", "No real national IDs in demos", "Internal", "Internal", "SYN-QA-/SYN-AE- formats", "Implemented in templates", "Master data", ""],
        ["CTL-08", "Audit", "Access logging", "Confirm", "Confirm", "Tryton + nginx logs", "Partial", "VM logs", "Central SIEM TBD"],
        ["CTL-09", "Tenancy", "Country/entity isolation", "Governance", "Governance", "company + record rules TBD", "Gap", "gap-analysis", "Cross-border leakage"],
        ["CTL-10", "Integrity", "Backup & restore", "Confirm", "Confirm", "PostgreSQL backups TBD", "Needs verification", "Ops", "Document RPO/RTO"],
        ["CTL-11", "Interop", "Secure HIE channels", "NABIDH/Riayati", "MoPH TBD", "Not implemented", "Gap", "interface-register", "No production HIE"],
        ["CTL-12", "Vendor", "No GH upgrade without CAB", "Internal", "Internal", "Policy", "Policy", "implementation guide", ""],
    ]
    add_rows(ws, headers, rows)
    path = SECURITY / "security-control-matrix.xlsx"
    wb.save(path)
    return path


def interface_register() -> Path:
    INTEGRATION.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Interfaces"
    headers = [
        "InterfaceID", "Name", "Direction", "Protocol", "Partner",
        "Country", "Status", "Owner", "Notes",
    ]
    rows = [
        ["IF-001", "React → Tryton JSON-RPC", "Outbound", "JSON-RPC/HTTPS", "GNU Health", "Both", "Live", "Engineering", "Primary app API"],
        ["IF-002", "SAO → Tryton", "Outbound", "JSON-RPC/HTTPS", "GNU Health", "Both", "Live", "Engineering", "Admin/wizards"],
        ["IF-003", "NABIDH ADT/ORU/etc.", "Outbound", "HL7v2", "DHA", "UAE-DXB", "Planned", "Integration", "Needs confirmation"],
        ["IF-004", "NABIDH C-CDA", "Outbound", "C-CDA", "DHA", "UAE-DXB", "Planned", "Integration", "Needs confirmation"],
        ["IF-005", "Malaffi clinical exchange", "Bidirectional", "TBD", "DoH", "UAE-AUH", "Planned", "Integration", "Confirm specs"],
        ["IF-006", "Riayati HIE", "Bidirectional", "HL7/FHIR/IHE", "MOHAP", "UAE", "Planned", "Integration", "Confirm profiles"],
        ["IF-007", "Qatar national e-health", "Bidirectional", "TBD", "MoPH", "QA", "Unknown", "Compliance", "Confirm existence/mandate"],
        ["IF-008", "Orthanc PACS", "Bidirectional", "DICOM", "Internal", "Both", "Optional", "Imaging", "health_orthanc"],
        ["IF-009", "Claims / eClaim", "Outbound", "TBD", "Payer/regulator", "UAE", "Not started", "Billing", "Needs confirmation"],
        ["IF-010", "SMTP / notifications", "Outbound", "SMTP", "Email gateway", "Both", "Not started", "Ops", "No PHI in subject lines"],
    ]
    add_rows(ws, headers, rows)
    path = INTEGRATION / "interface-register.xlsx"
    wb.save(path)
    return path


def master_data_workbook() -> Path:
    TEMPLATES.mkdir(parents=True, exist_ok=True)
    wb = Workbook()

    # Cover
    cover = wb.active
    cover.title = "Cover"
    cover["A1"] = "Healthcare Implementation Master Data"
    cover["A1"].font = Font(bold=True, size=16, color="1F4E79")
    cover["A3"] = "Product"
    cover["B3"] = "IST Health (GNU Health 5 / Tryton 7)"
    cover["A4"] = "Version"
    cover["B4"] = "1.0.0"
    cover["A5"] = "Date"
    cover["B5"] = TODAY
    cover["A6"] = "Classification"
    cover["B6"] = "SYNTHETIC / IMPLEMENTATION TEMPLATE — NO REAL PII"
    cover["A8"] = "synthetic_data"
    cover["B8"] = True
    cover["A10"] = "Instructions"
    cover["B10"] = "See Instructions sheet. Extensible taxonomies use ValidationLists. Freeze headers + filters enabled on data sheets."

    instructions = wb.create_sheet("Instructions")
    add_rows(
        instructions,
        ["Step", "Action"],
        [
            ["1", "Complete Cover metadata and VersionHistory"],
            ["2", "Fill sheets 6.1–6.10 using ValidationLists where dropdowns apply"],
            ["3", "Use only synthetic IDs: SYN-AE-########, SYN-QA-########, IST-P-#####"],
            ["4", "Set synthetic_data=TRUE on every demographic row"],
            ["5", "Run: python scripts/implementation/validate_master_data.py templates/Healthcare_Implementation_Master_Data.xlsx"],
            ["6", "Obtain SignOff before import to any shared environment"],
            ["7", "Never import to production without --allow-production override + CAB"],
        ],
    )

    vh = wb.create_sheet("VersionHistory")
    add_rows(vh, ["Version", "Date", "Author", "Changes"], [
        ["1.0.0", TODAY, "IST Framework Generator", "Initial template"],
    ])

    lists = wb.create_sheet("ValidationLists")
    add_rows(
        lists,
        ["ListName", "Value"],
        [
            ["Country", "AE"], ["Country", "QA"],
            ["EntityType", "Group"], ["EntityType", "LegalEntity"], ["EntityType", "Hospital"],
            ["EntityType", "Clinic"], ["EntityType", "Lab"], ["EntityType", "Pharmacy"],
            ["EntityType", "Warehouse"], ["EntityType", "Finance"],
            ["FacilityType", "Hospital"], ["FacilityType", "Clinic"], ["FacilityType", "Polyclinic"],
            ["FacilityType", "Diagnostic"], ["FacilityType", "Pharmacy"],
            ["Role", "SystemAdmin"], ["Role", "FacilityAdmin"], ["Role", "Physician"],
            ["Role", "Nurse"], ["Role", "Receptionist"], ["Role", "LabTech"],
            ["Role", "Pharmacist"], ["Role", "Billing"], ["Role", "Warehouse"],
            ["Role", "ImagingTech"], ["Role", "Compliance"], ["Role", "Executive"],
            ["Role", "ITSupport"], ["Role", "MedicalRecords"],
            ["YesNo", "TRUE"], ["YesNo", "FALSE"],
            ["Env", "dev"], ["Env", "test"], ["Env", "uat"], ["Env", "prod"],
        ],
    )

    sheets_spec = [
        ("6.1_LegalEntities", ["EntityCode", "EntityName", "EntityName_AR", "Country", "EntityType", "ParentCode", "TaxID_Synthetic", "Currency", "Active", "synthetic_data", "Notes"]),
        ("6.2_Facilities", ["FacilityCode", "FacilityName", "FacilityName_AR", "Country", "EmirateOrMunicipality", "FacilityType", "ParentEntityCode", "Address_Synthetic", "Phone_Synthetic", "Active", "synthetic_data"]),
        ("6.3_Departments", ["DeptCode", "DeptName", "DeptName_AR", "FacilityCode", "SpecialtyCode", "Active", "synthetic_data"]),
        ("6.4_UsersRoles", ["Login", "DisplayName", "Role", "FacilityCode", "Country", "Language", "Active", "synthetic_data", "TempPasswordFlag"]),
        ("6.5_Specialties", ["SpecialtyCode", "SpecialtyName", "SpecialtyName_AR", "ServiceLine", "Active", "ExtensibleTag"]),
        ("6.6_ProductsMeds", ["ProductCode", "Name", "Name_AR", "Category", "Unit", "ListPrice", "Active", "synthetic_data"]),
        ("6.7_LabCatalog", ["TestCode", "TestName", "Specimen", "LOINC_Optional", "Active", "Notes"]),
        ("6.8_Insurance", ["PlanCode", "PayerName", "Country", "CoverageType", "Active", "synthetic_data", "NeedsLegalConfirm"]),
        ("6.9_Warehouses", ["LocationCode", "LocationName", "FacilityCode", "LocationType", "Active", "synthetic_data"]),
        ("6.10_Sequences", ["SequenceCode", "Purpose", "Prefix", "Padding", "NextNumber", "Country", "Notes"]),
    ]

    sample_rows = {
        "6.1_LegalEntities": [
            ["IST-GRP", "IST Health Group", "مجموعة آي إس تي الصحية", "AE", "Group", "", "SYN-TAX-0001", "AED", "TRUE", "TRUE", "Holding"],
            ["IST-AE", "IST Health UAE LLC", "آي إس تي الإمارات", "AE", "LegalEntity", "IST-GRP", "SYN-TAX-0002", "AED", "TRUE", "TRUE", ""],
            ["IST-QA", "IST Health Qatar LLC", "آي إس تي قطر", "QA", "LegalEntity", "IST-GRP", "SYN-TAX-0003", "QAR", "TRUE", "TRUE", ""],
        ],
        "6.2_Facilities": [
            ["FAC-DXB-01", "IST Dubai Clinic", "عيادة دبي", "AE", "Dubai", "Clinic", "IST-AE", "SYN Address Dubai", "+971-50-000-0001", "TRUE", "TRUE"],
            ["FAC-DOH-01", "IST Doha Clinic", "عيادة الدوحة", "QA", "Doha", "Clinic", "IST-QA", "SYN Address Doha", "+974-3000-0001", "TRUE", "TRUE"],
        ],
        "6.3_Departments": [
            ["DEPT-GM", "General Medicine", "الطب العام", "FAC-DOH-01", "SPEC-GP", "TRUE", "TRUE"],
            ["DEPT-LAB", "Laboratory", "المختبر", "FAC-DOH-01", "SPEC-LAB", "TRUE", "TRUE"],
        ],
        "6.4_UsersRoles": [
            ["syn.admin", "Synthetic Admin", "SystemAdmin", "FAC-DOH-01", "QA", "en", "TRUE", "TRUE", "TRUE"],
            ["syn.physician", "Synthetic Physician", "Physician", "FAC-DOH-01", "QA", "ar", "TRUE", "TRUE", "TRUE"],
        ],
        "6.5_Specialties": [
            ["SPEC-GP", "General Practice", "الطب العام", "Outpatient", "TRUE", "extensible"],
            ["SPEC-LAB", "Laboratory Medicine", "طب المختبر", "Diagnostics", "TRUE", "extensible"],
        ],
        "6.6_ProductsMeds": [
            ["IST-MED-0001", "Paracetamol 500mg", "باراسيتامول", "Medicaments", "tab", 5, "TRUE", "TRUE"],
        ],
        "6.7_LabCatalog": [
            ["LAB-CBC", "Complete Blood Count", "Blood", "", "TRUE", "LOINC optional pending license"],
        ],
        "6.8_Insurance": [
            ["INS-SYN-01", "Synthetic Payer A", "QA", "Outpatient", "TRUE", "TRUE", "TRUE"],
        ],
        "6.9_Warehouses": [
            ["WH-PHARM", "Pharmacy Store", "FAC-DOH-01", "Pharmacy", "TRUE", "TRUE"],
        ],
        "6.10_Sequences": [
            ["SEQ-PAT-QA", "Patient code QA", "IST-P-", 5, 1, "QA", "Synthetic"],
            ["SEQ-PAT-AE", "Patient code AE", "IST-AE-P-", 5, 1, "AE", "Synthetic"],
        ],
    }

    for name, headers in sheets_spec:
        ws = wb.create_sheet(name[:31])
        ws.append(headers)
        for row in sample_rows.get(name, []):
            ws.append(row)
        style_header(ws, len(headers))
        autosize(ws)
        # Yes/No validation on synthetic_data columns if present
        if "synthetic_data" in headers:
            idx = headers.index("synthetic_data") + 1
            letter = get_column_letter(idx)
            dv = DataValidation(type="list", formula1='"TRUE,FALSE"', allow_blank=True)
            ws.add_data_validation(dv)
            dv.add(f"{letter}2:{letter}5000")

    dd = wb.create_sheet("DataDictionary")
    add_rows(
        dd,
        ["Sheet", "Field", "Type", "Required", "Description"],
        [
            ["6.1_LegalEntities", "EntityCode", "string", "Y", "Stable code; no real trade license numbers"],
            ["6.2_Facilities", "Phone_Synthetic", "string", "N", "Synthetic phone only"],
            ["6.4_UsersRoles", "Login", "string", "Y", "Non-PII synthetic login"],
            ["All", "synthetic_data", "boolean", "Y", "Must be TRUE for non-production rows"],
            ["6.7_LabCatalog", "LOINC_Optional", "string", "N", "Populate only with valid LOINC license"],
        ],
    )

    sign = wb.create_sheet("SignOff")
    add_rows(
        sign,
        ["Role", "Name", "Date", "Signature", "Decision"],
        [
            ["Clinical Lead", "", "", "", ""],
            ["Compliance/Legal", "", "", "", ""],
            ["IT Security", "", "", "", ""],
            ["Implementation Lead", "", "", "", ""],
        ],
    )

    path = TEMPLATES / "Healthcare_Implementation_Master_Data.xlsx"
    wb.save(path)
    return path


def main() -> None:
    paths = [
        regulatory_matrix(),
        capability_matrix(),
        security_matrix(),
        interface_register(),
        master_data_workbook(),
    ]
    for p in paths:
        print(f"Wrote {p} exists={p.exists()} size={p.stat().st_size}")


if __name__ == "__main__":
    main()
