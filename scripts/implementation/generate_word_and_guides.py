#!/usr/bin/env python3
"""Generate Word (.docx) and major Markdown implementation / user / security docs."""

from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.shared import Pt

ROOT = Path(__file__).resolve().parents[2]
TODAY = date.today().isoformat()


def md(path: Path, body: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(body.strip() + "\n", encoding="utf-8")
    print(f"Wrote {path}")


def docx_from_sections(path: Path, title: str, sections: list[tuple[str, str]]) -> None:
    doc = Document()
    doc.add_heading(title, 0)
    p = doc.add_paragraph(f"Generated: {TODAY} | IST Health Implementation Framework")
    p.runs[0].font.size = Pt(10)
    for heading, text in sections:
        doc.add_heading(heading, level=1)
        for para in text.strip().split("\n\n"):
            doc.add_paragraph(para.strip())
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)
    print(f"Wrote {path}")


def implementation_guide() -> None:
    phases = "\n".join(
        f"### Phase {i}\n{desc}\n"
        for i, desc in enumerate(
            [
                "Discovery & baseline (architecture, modules, VM verification).",
                "Regulatory & privacy kickoff (UAE/QA counsel; classify needs-confirmation items).",
                "Org model design: Group → UAE/QA legal entities → facilities.",
                "Environment strategy (dev/test/uat/prod) and synthetic-data gates.",
                "Tryton company/institution/units configuration (no GH upgrade).",
                "Master data collection via Excel template + questionnaire.",
                "Role packs & access matrix (res.group / record rules).",
                "Clinical catalog: specialties, products, lab types, sequences.",
                "UAE-specific config (emirate, HIE readiness flags) — pending authority confirmation.",
                "Qatar-specific config (MoPH facility codes) — pending confirmation.",
                "Arabic/EN localization & RTL QA.",
                "Synthetic data load (deterministic, idempotent) in non-prod.",
                "Integration design (HL7/FHIR adapter) — no production HIE without approval.",
                "Security controls (TLS, backups, audit, session hardening).",
                "UAT with role-based scenarios + Help Centre guided workflows.",
                "Training & user guides sign-off.",
                "Cutover rehearsal / rollback plan.",
                "Production import (CAB + --allow-production).",
                "Hypercare (30 days).",
                "HIE onboarding (if mandated) with partner.",
                "Continuous compliance review (PDPL/PDPPL).",
                "Optimization & reporting handover.",
            ],
            start=1,
        )
    )

    body = f"""
# UAE & Qatar Implementation Guide — IST Health

**Version:** 1.0.0  
**Date:** {TODAY}  
**Stack:** GNU Health HIS 5 / Tryton 7 + IST React frontend  
**Rule:** Do not upgrade GNU Health/Tryton in this program without CAB.

## 1. Purpose

Deliver a production-quality implementation framework for multi-country (UAE + Qatar) healthcare operations on the existing IST Health platform, with governance, master data, synthetic factories, help, and security artefacts.

## 2. Principles

- Prefer reusable services (integration, identity, audit) over one-off scripts.
- Never use real Emirates ID, Qatar ID, passports, or real phones/emails in non-prod.
- Mark all synthetic rows (`synthetic_data=true` or code prefix markers).
- Do not fabricate regulatory requirements; uncertain items require legal/authority confirmation.
- Do not invent GNU Health capabilities; verify against docs/VM.

## 3. Organization model & tenant isolation

```
IST Health Group
├── Legal Entity — UAE (AED)
│   ├── Hospitals / Clinics / Labs / Pharmacies / Warehouses / Finance
│   └── Emirate attributes (Dubai / Abu Dhabi / Northern)
└── Legal Entity — Qatar (QAR)
    └── Hospitals / Clinics / Labs / Pharmacies / Warehouses / Finance
```

**Isolation approach (proposed):**

| Layer | Mechanism | Status |
|-------|-----------|--------|
| Legal entity | `company.company` per country | To configure |
| Facility | `gnuhealth.institution` (+ units) | Partial (single demo) |
| Data | Tryton record rules by company/institution | To design |
| UI | Country/role context in Help + future filters | Help awareness added |
| DB | Prefer single DB multi-company first; separate DB only if mandated | Confirm with security |

Cross-entity clinical chart sharing must be an explicit clinical/legal decision.

## 4. UAE vs Qatar configuration differences

| Topic | UAE | Qatar |
|-------|-----|-------|
| Currency | AED | QAR |
| HIE | NABIDH / Malaffi / Riayati — **confirm per facility** | MoPH e-health — **confirm** |
| National ID | Emirates ID rules — **confirm**; demos use `SYN-AE-########` | Qatar ID — **confirm**; demos use `SYN-QA-########` |
| Language | EN + AR | EN + AR |
| Privacy law | UAE PDPL — **confirm scope** | PDPPL Law 13/2016 — **confirm** |
| Claims | eClaim/schemes — **confirm** | Local insurance — **confirm** |

## 5. Phases 1–22

{phases}

## 6. Governance

### RACI (summary)

| Activity | Sponsor | Compliance | Clinical | Engineering | Ops |
|----------|---------|------------|----------|-------------|-----|
| Regulatory interpretation | A | R | C | I | I |
| Org / tenant model | A | C | C | R | C |
| Master data quality | C | C | A/R | C | I |
| Synthetic loads | I | C | C | R | A |
| HIE go-live | A | R | C | R | C |
| Production import | A | R | C | R | R |

### Risks

| Risk | Impact | Mitigation |
|------|--------|------------|
| Assuming HIE mandates without confirmation | Legal/ops | Source register + counsel |
| Cross-country data leakage | High | Record rules + company split |
| Real PII in demos | High | Import gates + synthetic markers |
| GH upgrade drift | High | Freeze version; CAB only |
| localStorage session secrets | Medium | Phase-2 auth hardening |

## 7. Master data & questionnaire

- Excel: `templates/Healthcare_Implementation_Master_Data.xlsx` (sheets 6.1–6.10)
- Questionnaire: `templates/Healthcare_Implementation_Questionnaire.docx`
- Validate: `python scripts/implementation/validate_master_data.py …`

## 8. Synthetic data

See `docs/implementation/synthetic-data-guide.md` and `data-import-runbook.md`.  
Demo config: `config/synthetic/uae_qatar_demo.yaml`.

## 9. Help & training

React Help Centre loads Markdown/JSON from `frontend/help-content/{{en|ar}}/`.  
Role guides: `docs/user-guides/`.

## 10. Exit criteria (go-live)

- [ ] Legal confirmation logged for open regulatory items or risk acceptance
- [ ] Tenant isolation test passed
- [ ] Synthetic gate blocks prod without override
- [ ] UAT sign-off by role
- [ ] Backup/restore drill evidenced
- [ ] No unapproved GH/Tryton upgrade
"""
    md(ROOT / "docs" / "implementation" / "UAE_Qatar_Implementation_Guide.md", body)
    docx_from_sections(
        ROOT / "docs" / "implementation" / "UAE_Qatar_Implementation_Guide.docx",
        "UAE & Qatar Implementation Guide — IST Health",
        [
            ("Purpose", "Production-quality multi-country implementation framework for IST Health on GNU Health 5 / Tryton 7."),
            ("Principles", "Reusable services; synthetic-only non-prod; no fabricated regulations; no invented GH capabilities; no unapproved upgrades."),
            ("Org model", "Group → UAE/QA legal entities → hospitals/clinics/labs/pharmacies/warehouses/finance with Tryton company + institution isolation."),
            ("Phases", "Execute phases 1–22 from the Markdown guide: discovery through hypercare, HIE onboarding, and continuous compliance."),
            ("Governance", "RACI across Sponsor, Compliance, Clinical, Engineering, Ops. Maintain risk register for HIE assumptions and PII leakage."),
            ("References", "See docs/research/*, templates/*, docs/security/*, docs/integration/*."),
        ],
    )


def questionnaire() -> None:
    docx_from_sections(
        ROOT / "templates" / "Healthcare_Implementation_Questionnaire.docx",
        "Healthcare Implementation Questionnaire — UAE / Qatar",
        [
            ("A. Organization", "Legal entities, trade names (synthetic OK for pilots), countries of operation, primary facilities, reporting currency."),
            ("B. Facilities", "Hospitals/clinics/labs/pharmacies counts; emirate/municipality; bed capacity; operating hours; bilingual requirements."),
            ("C. Clinical scope", "Outpatient, inpatient, ED, surgery, ICU, maternity, pediatrics, dental, ophthalmology — in/out of scope."),
            ("D. Laboratory & imaging", "In-house vs outsourced; analyzers; PACS/Orthanc interest; worklist needs."),
            ("E. Pharmacy & stock", "Formulary source; narcotics controls; warehouse topology."),
            ("F. Users & roles", "Approximate counts per role; shift patterns; shared workstation policy."),
            ("G. Identity & privacy", "Will national IDs be stored? Consent model? Retention period? Cross-border access? — legal confirmation required."),
            ("H. Interoperability", "NABIDH / Malaffi / Riayati / Qatar HIE obligations known? HL7/FHIR/DICOM already in use?"),
            ("I. Billing & insurance", "Payers; claims platforms; cash vs insurance mix — confirm regulatory claims interfaces."),
            ("J. Non-functional", "Expected concurrent users; RPO/RTO; languages; mobile needs; reporting."),
            ("K. Environments", "Who may access prod? Synthetic data policy acknowledgement."),
            ("Sign-off", "Completed by / date / clinical lead / compliance / IT."),
        ],
    )


def security_md() -> None:
    md(
        ROOT / "docs" / "security" / "uae-qatar-health-data-controls.md",
        f"""
# UAE / Qatar Health Data Controls

**Date:** {TODAY}  
**Disclaimer:** Control mapping is engineering guidance. Legal applicability of UAE PDPL and Qatar PDPPL to specific processing activities **requires confirmation** from qualified counsel and competent authorities.

## 1. Data classification

| Class | Examples | Handling |
|-------|----------|----------|
| Synthetic demo | `IST-P-*`, `SYN-AE-*`, `SYN-QA-*` | Allowed in non-prod; mark synthetic |
| Real PHI/PII | Real names, national IDs, phones | Prod only; least privilege; audit |
| Special nature (QA PDPPL) | Health data | Enhanced safeguards — **confirm permits** |
| Sensitive (UAE PDPL) | Health/biometric | **Confirm** lawful basis & exemptions |

## 2. Technical controls (current vs target)

| Control | Current IST | Target |
|---------|-------------|--------|
| TLS | HTTPS via nginx (verify) | TLS1.2+ only; HSTS |
| Auth session | localStorage MVP | httpOnly cookie / JWT gateway |
| RBAC | Tryton groups | Named IST role packs |
| Tenant isolation | Single company demo | Company + institution record rules |
| Encryption at rest | GCP disk default? | Verify CMEK if required |
| Backup | TBD on VM | Documented restore drill |
| Synthetic gate | In progress | Refuse prod import w/o override |
| Logging | nginx + trytond | Central audit for PHI reads |

## 3. Operational controls

1. Non-production databases must contain only synthetic data.  
2. Production access via named accounts; no shared admin passwords in chat.  
3. DPIA before HIE go-live or new secondary use.  
4. Breach response contacts (UAE/QA) — **confirm notification timelines with counsel**.  
5. Vendor (GH/Tryton) upgrades only via CAB.

## 4. Cross-border

UAE ↔ Qatar chart sharing is **not enabled by default**. Any cross-entity access needs legal basis + technical rule.

## 5. Related artefacts

- `security-control-matrix.xlsx`
- `docs/integration/interoperability-assessment.md`
""",
    )


def interoperability_md() -> None:
    md(
        ROOT / "docs" / "integration" / "interoperability-assessment.md",
        f"""
# Interoperability Assessment — IST Health

**Date:** {TODAY}

## Current state

| Interface | Status |
|-----------|--------|
| React ↔ Tryton JSON-RPC | Live |
| SAO ↔ Tryton | Live |
| HL7 v2 | Not implemented |
| FHIR | Not implemented |
| C-CDA | Not implemented |
| DICOM / Orthanc | Package may exist; not wired in React |
| NABIDH / Malaffi / Riayati | Not implemented |
| Qatar national HIE | Mandate/spec **unconfirmed** |

## Standards notes (licensing)

| Standard | Note |
|----------|------|
| HL7 / FHIR | Implement freely per HL7 policies; respect trademarks |
| DICOM | Free standard; follow NEMA terms |
| ICD | Confirm national/commercial licensing |
| SNOMED CT | Confirm affiliate license for UAE & Qatar |
| LOINC | Free with license attribution |

## Recommended design (reusable service)

```
HIS (Tryton) → Outbox / events → Integration Service → HL7/FHIR adapters → HIE endpoints
```

Do **not** embed HIE credentials in the React app. Prefer a server-side integration service with audit, retries, and payload minimization.

## Next actions

1. Confirm which facilities are legally required to connect to which HIE.  
2. Obtain current interface specs from DHA/DoH/MOHAP/MoPH.  
3. Prototype ADT^A01 + ORU^R01 against a vendor test harness (non-prod).  
4. Update `interface-register.xlsx` statuses.
""",
    )


def data_import_runbook() -> None:
    md(
        ROOT / "docs" / "implementation" / "data-import-runbook.md",
        """
# Data Import Runbook

## Preconditions

- Target environment classified (dev/test/uat/prod)
- Master data workbook validated
- Backup taken for shared environments
- For prod: CAB approval + `--allow-production`

## Commands

```bash
# Validate master data Excel
python scripts/implementation/validate_master_data.py templates/Healthcare_Implementation_Master_Data.xlsx

# Validate synthetic config
python scripts/synthetic_data/cli.py validate --config config/synthetic/uae_qatar_demo.yaml

# Dry-run import (no writes)
python scripts/synthetic_data/cli.py import_data --config config/synthetic/uae_qatar_demo.yaml --dry-run

# Generate synthetic dataset (local files / report)
python scripts/synthetic_data/cli.py generate_synthetic --config config/synthetic/uae_qatar_demo.yaml

# Import synthetic into GNU Health (non-prod)
python scripts/synthetic_data/cli.py import_synthetic --config config/synthetic/uae_qatar_demo.yaml

# Legacy engine (existing)
python run_synthetic_data.py --config config.yaml --dry-run
```

## Rollback

1. Stop imports.  
2. Restore PostgreSQL backup taken pre-import.  
3. Or delete by synthetic code prefix (`IST-P-*`, comments containing `IST synthetic`) — only if documented and tested.  
4. Record incident in import report under `docs/testing/`.

## Refuse production

CLI exits non-zero if `environment: production` (or server URL matches prod host) unless `--allow-production` is passed.
""",
    )


def synthetic_guide() -> None:
    md(
        ROOT / "docs" / "implementation" / "synthetic-data-guide.md",
        """
# Synthetic Data Guide

## Goals

Deterministic, idempotent, bilingual-capable demo data for UAE + Qatar **without real PII**.

## Markers

| Mechanism | Example |
|-----------|---------|
| Code prefix | `IST-P-00001`, `IST-ST-DOC-0001` |
| Comment/text | `IST synthetic: …` |
| Config flag | `synthetic_data: true` |
| National ID (fake) | `SYN-QA-00000001`, `SYN-AE-00000001` |

## Determinism

`seed: 42` (or config value) seeds Python `random` and Faker.

## Scale

| Profile | Patients | Staff | Use |
|---------|----------|-------|-----|
| Demo (`uae_qatar_demo.yaml`) | 50 | 12 | CI / evidence |
| Default `config.yaml` | 4000 | 300 | VM load test |
| Full brief optional | 50000 | — | Only on sized VM; document runtime |

## Safety

- Default idempotent top-up  
- `--dry-run` validates connectivity / plan  
- Production refused without override  
""",
    )


# Section 10 — all roles required by the master brief
ROLES = [
    ("group-admin", "Group Administrator", "Group-level companies, multi-country entities, cross-facility governance"),
    ("country-admin", "Country Administrator", "UAE or Qatar legal entity settings, currency, country-scoped users"),
    ("facility-admin", "Facility Administrator", "Institution settings, departments, local users"),
    ("system-admin", "System Administrator", "Users, groups, modules, sequences, companies, settings"),
    ("receptionist", "Registration and Reception", "Patients create, appointments, check-in, demographics"),
    ("call-centre", "Call Centre", "Phone booking, appointment reminders, patient inquiries"),
    ("nurse", "Nurse", "Nursing, beds, ambulatory care, vitals support workflows"),
    ("triage-nurse", "Triage Nurse", "ED/OPD triage acuity, initial vitals, routing to clinician"),
    ("physician", "Doctor", "Patients, evaluations, prescriptions, appointments, procedures"),
    ("surgeon", "Surgeon", "Surgery requests, OR scheduling, operative notes (SAO deep)"),
    ("lab-tech", "Laboratory Technician", "Laboratory sample processing and results entry"),
    ("lab-supervisor", "Laboratory Supervisor", "Lab result approval, QC, analyzer configuration"),
    ("imaging-tech", "Radiology Technician", "Imaging requests, acquisition, worklist"),
    ("radiologist", "Radiologist", "Imaging report interpretation and approval"),
    ("pharmacist", "Pharmacist", "Prescription review, dispensing, controlled drugs"),
    ("pharmacy-assistant", "Pharmacy Assistant", "Dispense support, stock picks, labelling under pharmacist oversight"),
    ("cashier", "Cashier", "Patient payments, receipts, cash session close"),
    ("billing", "Billing Employee", "Invoices, products, patient billing queues"),
    ("insurance-claims", "Insurance and Claims Employee", "Payer claims, eClaim preparation, rejection handling"),
    ("finance", "Finance Employee", "Chart of accounts, periods, journals (primarily SAO)"),
    ("procurement", "Procurement Employee", "Purchase orders, supplier parties, receiving coordination"),
    ("warehouse", "Inventory Employee", "Stock locations, moves, inventory counts"),
    ("hr", "HR Employee", "Health professionals, party employment attributes, scheduling support"),
    ("dept-manager", "Department Manager", "Department KPIs, staff roster awareness, local approvals"),
    ("executive", "Hospital Management", "Dashboard KPIs, executive reporting (SAO)"),
    ("it-support", "IT Support", "Login troubleshooting, HTTPS, module activation ops"),
    ("auditor", "Auditor", "Read-only audit of access, synthetic policy, compliance evidence packs"),
]


def user_guides() -> None:
    role_links = "\n".join(f"- [{title}](roles/{slug}.md)" for slug, title, _ in ROLES)
    md(
        ROOT / "docs" / "user-guides" / "Role_Based_User_Guide.md",
        f"""
# Role-Based User Guide — IST Health

**Date:** {TODAY}  
**App:** React UI at `/` with SAO fallback for deep wizards.

## Roles

{role_links}

## Common tasks

| Task | Route |
|------|-------|
| Sign in | `/login` |
| Dashboard | `/` |
| Patients | `/patients` |
| Appointments | `/appointments` |
| Help Centre | `/help` |
| Language | Header EN/AR toggle (RTL for AR) |

## SAO

Use **Open in SAO** for chart of accounts, complex wizards, and unmapped menus.
""",
    )
    docx_from_sections(
        ROOT / "docs" / "user-guides" / "Role_Based_User_Guide.docx",
        "Role-Based User Guide — IST Health",
        [("Overview", "IST Health React UI over GNU Health. Use role sheets for task lists. Arabic via language switch.")]
        + [(title, f"Primary focus: {focus}. See Markdown role file {slug}.md for routes.") for slug, title, focus in ROLES],
    )
    for slug, title, focus in ROLES:
        md(
            ROOT / "docs" / "user-guides" / "roles" / f"{slug}.md",
            f"""
# {title}

**Focus:** {focus}

## Access

Assigned via Tryton `res.group`. React currently shows menus based on Tryton menu ACL.

## Typical routes

See main Role_Based_User_Guide.md route table; prioritize screens listed in focus.

## Guided help

Open Help Centre → workflows for this role (EN/AR).

## Notes

- Do not enter real national IDs in training databases.
- Prefer synthetic patients (`IST-P-*`).
""",
        )


def main() -> None:
    implementation_guide()
    questionnaire()
    security_md()
    interoperability_md()
    data_import_runbook()
    synthetic_guide()
    user_guides()
    print("Markdown/DOCX generation complete.")


if __name__ == "__main__":
    main()
